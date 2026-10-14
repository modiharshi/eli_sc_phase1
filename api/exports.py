"""Text exports of judgments (txt, docx, pdf). Original PDFs come in Stage 10."""
import io

from fastapi import HTTPException
from fastapi.responses import Response

from . import db
from .common import ATTRIBUTION, NOT_OFFICIAL

NOTICE = f"{NOT_OFFICIAL} {ATTRIBUTION}"


def _text(j, lang):
    t = db.q1("select text from judgment_text where judgment_id = %s and language_code = %s", (j["judgment_id"], lang))
    return t["text"] if t else None


def _head(j):
    return [("Title", j["title"]), ("Decision date", j["decision_date"]), ("INSC id", j["insc_id"]), ("Citation", j["citation"]),
            ("Bench strength", j["bench_strength"]), ("Disposal", j["disposal"]), ("Permalink", f"/j/{j['judgment_id']}")]


def single(j, fmt, lang="en"):
    text = _text(j, lang)
    if text is None:
        raise HTTPException(404, f"This judgment has no {lang} text to export (metadata only).")
    return bundle([j], fmt, {j["judgment_id"]: text})


def bundle(js, fmt, texts=None):
    texts = texts or {j["judgment_id"]: _text(j, "en") for j in js}
    texts = {k: v for k, v in texts.items() if v}
    if not texts:
        raise HTTPException(404, "None of the selected judgments has English text.")
    js = [j for j in js if j["judgment_id"] in texts]
    name = "eli_judgment" if len(js) == 1 else "eli_judgments"
    if fmt == "txt":
        out = "\n\n".join("\n".join(f"{k}: {v}" for k, v in _head(j) if v) + "\n\n" + texts[j["judgment_id"]] for j in js) + f"\n\n---\n{NOTICE}\n"
        return Response(out, media_type="text/plain; charset=utf-8", headers={"Content-Disposition": f"attachment; filename={name}.txt"})
    if fmt == "docx":
        from docx import Document
        d = Document()
        d.add_paragraph(NOT_OFFICIAL).runs[0].bold = True
        for j in js:
            d.add_heading(j["title"] or "(untitled)", level=1)
            for k, v in _head(j):
                if v:
                    d.add_paragraph(f"{k}: {v}")
            for para in texts[j["judgment_id"]].split("\n\n"):
                if para.strip():
                    d.add_paragraph(para.strip())
            d.add_page_break()
        d.add_paragraph(ATTRIBUTION)
        buf = io.BytesIO()
        d.save(buf)
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": f"attachment; filename={name}.docx"})
    if fmt == "pdf":
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape

        def clean(s):  # base-14 PDF fonts are Latin-1 only; unsupported characters are replaced, the text export keeps them
            return escape(str(s)).encode("latin-1", "replace").decode("latin-1")
        st = getSampleStyleSheet()
        buf = io.BytesIO()
        story = [Paragraph(f"<b>{clean(NOT_OFFICIAL)}</b>", st["Normal"]), Spacer(1, 8)]
        for j in js:
            story.append(Paragraph(clean(j["title"] or "(untitled)"), st["Heading2"]))
            for k, v in _head(j):
                if v:
                    story.append(Paragraph(clean(f"{k}: {v}"), st["Normal"]))
            story.append(Spacer(1, 8))
            for para in texts[j["judgment_id"]].split("\n\n"):
                if para.strip():
                    story.append(Paragraph(clean(para.strip()).replace("\n", "<br/>"), st["BodyText"]))
            story.append(PageBreak())
        story.append(Paragraph(clean(ATTRIBUTION), st["Normal"]))
        SimpleDocTemplate(buf, pagesize=A4, title="ELI export").build(story)
        return Response(buf.getvalue(), media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={name}.pdf"})
    raise HTTPException(400, "format must be txt, docx or pdf")
